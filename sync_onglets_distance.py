#!/usr/bin/env python3
"""Regenere les onglets 10K / 21K / 42K a partir de l'onglet ALL.

Ces onglets ne sont lus par AUCUN script : ils servent uniquement de vue par
distance quand on ouvre le fichier a la main. Livres a eux-memes ils derivent —
49 divergences avec ALL et 768 cellules manquantes avaient ete constatees en
septembre 2026. Ce script en fait un miroir exact de ALL.

A relancer apres toute modification des finishers :
    python sync_onglets_distance.py
"""
import sys
from pathlib import Path

import openpyxl

FICHIER = Path(__file__).parent / "Suivi_Finishers_Monde_10k_-_21k_-_42k_HISTORIQUE.xlsx"
ONGLETS = {"10K": "10KM", "21K": "SEMI", "42K": "MARATHON"}


def main():
    wb = openpyxl.load_workbook(FICHIER)
    src = wb["ALL"]
    hdr = [c.value for c in src[1]]
    i_per, i_city, i_dist, i_race = (hdr.index("Période"), hdr.index("City"),
                                     hdr.index("Distance"), hdr.index("Race"))
    annees = [(i, h) for i, h in enumerate(hdr) if isinstance(h, (int, float))]

    lignes = [r for r in src.iter_rows(min_row=2, values_only=True) if r[i_race]]
    total = 0
    for onglet, distance in ONGLETS.items():
        if onglet not in wb.sheetnames:
            print(f"  onglet {onglet} absent, ignore")
            continue
        ws = wb[onglet]
        wb.remove(ws)
        ws = wb.create_sheet(onglet)
        ws.append(["Période", "City", "Race"] + [int(a) for _, a in annees])
        n = 0
        for r in lignes:
            if str(r[i_dist] or "").strip() != distance:
                continue
            ws.append([r[i_per], r[i_city], r[i_race]] + [r[i] for i, _ in annees])
            n += 1
        print(f"  {onglet:4s} <- {n} lignes de ALL (distance {distance})")
        total += n

    # remettre les onglets dans l'ordre d'origine
    ordre = ["10K", "21K", "42K", "ALL", "BIGGEST EVENTS"]
    wb._sheets.sort(key=lambda s: ordre.index(s.title) if s.title in ordre else 99)
    wb.save(FICHIER)
    print(f"{total} lignes ecrites. Les onglets par distance refletent ALL.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
